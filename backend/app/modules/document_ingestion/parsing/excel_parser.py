from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


@dataclass(slots=True)
class ExcelSheet:
    name: str
    rows: list[list[Any]]


@dataclass(slots=True)
class ExcelWorkbook:
    sheets: list[ExcelSheet]
    metadata: dict[str, Any]


class ExcelDocumentParser:
    def parse(self, data: bytes, filename: str) -> ExcelWorkbook:
        if not filename.lower().endswith(".xlsx"):
            raise RuntimeError("job_spreadsheet_xlsx_required")
        try:
            workbook = load_workbook(
                BytesIO(data),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise RuntimeError("job_spreadsheet_invalid") from exc

        try:
            sheets: list[ExcelSheet] = []
            for worksheet in workbook.worksheets:
                if worksheet.sheet_state != "visible":
                    continue
                rows = [
                    _trim_row(list(row))
                    for row in worksheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(_cell_text(cell) for cell in row)]
                if rows:
                    sheets.append(ExcelSheet(name=worksheet.title, rows=rows))
            if not sheets:
                raise RuntimeError("job_spreadsheet_contains_no_rows")
            return ExcelWorkbook(
                sheets=sheets,
                metadata={
                    "provider": "openpyxl",
                    "format": "xlsx",
                    "sheet_names": [sheet.name for sheet in sheets],
                    "sheet_count": len(sheets),
                },
            )
        finally:
            workbook.close()


def _trim_row(row: list[Any]) -> list[Any]:
    while row and row[-1] is None:
        row.pop()
    return row


def _cell_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""
