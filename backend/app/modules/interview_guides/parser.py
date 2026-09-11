from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import fitz
from openpyxl import load_workbook


HEADER_ALIASES = {
    "question": {"题目", "问题", "面试题", "question"},
    "evaluation": {"考察要点", "评价要点", "参考要点", "evaluationpoints", "evaluation"},
    "required": {"是否必问", "必问", "required"},
    "result_type": {"结果类型", "记录类型", "resulttype"},
}


def parse_template_file(filename: str, data: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(data)
    if suffix == ".docx":
        return _parse_docx(data)
    if suffix == ".pdf":
        return _parse_pdf(data)
    raise ValueError("仅支持 .xlsx、.docx 和 .pdf 题单")


def _normalized_header(value: object) -> str:
    return re.sub(r"[\s_\-]+", "", str(value or "").strip().lower())


def _find_column(headers: list[object], key: str) -> int | None:
    aliases = HEADER_ALIASES[key]
    for index, value in enumerate(headers):
        if _normalized_header(value) in aliases:
            return index
    return None


def _split_points(value: object) -> list[str]:
    return [item.strip() for item in re.split(r"[\n；;、]+", str(value or "")) if item.strip()]


def _required(value: object) -> bool:
    return _normalized_header(value) not in {"否", "false", "no", "0", "非必问"}


def _result_type(value: object) -> str:
    normalized = _normalized_header(value)
    return "non_scoring" if normalized in {"非评分", "非评分信息", "不计分", "nonscoring"} else "capability"


def _parse_xlsx(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    question_column = _find_column(list(rows[0]), "question")
    if question_column is None:
        raise ValueError("Excel缺少“题目”列")
    evaluation_column = _find_column(list(rows[0]), "evaluation")
    required_column = _find_column(list(rows[0]), "required")
    result_type_column = _find_column(list(rows[0]), "result_type")
    questions: list[dict] = []
    for row in rows[1:]:
        question = str(row[question_column] or "").strip()
        if not question:
            continue
        questions.append({
            "question": question,
            "evaluationPoints": _split_points(row[evaluation_column]) if evaluation_column is not None and evaluation_column < len(row) else [],
            "required": _required(row[required_column]) if required_column is not None and required_column < len(row) else True,
            "resultType": _result_type(row[result_type_column]) if result_type_column is not None and result_type_column < len(row) else "capability",
        })
    return questions


def _parse_docx(data: bytes) -> list[dict]:
    try:
        with ZipFile(BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
    except (BadZipFile, KeyError) as exc:
        raise ValueError("Word文件结构无效") from exc
    root = ElementTree.fromstring(xml)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    lines: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            lines.append(text)
    return _questions_from_lines(lines)


def _parse_pdf(data: bytes) -> list[dict]:
    document = fitz.open(stream=data, filetype="pdf")
    lines: list[str] = []
    for page in document:
        lines.extend(line.strip() for line in page.get_text("text").splitlines() if line.strip())
    return _questions_from_lines(lines)


def _questions_from_lines(lines: list[str]) -> list[dict]:
    questions: list[dict] = []
    numbered = re.compile(r"^\s*(?:第\s*)?(\d+)[.、）)]\s*(.+)$")
    for raw in lines:
        line = re.sub(r"\s+", " ", raw).strip()
        match = numbered.match(line)
        if match:
            questions.append(_plain_question(match.group(2)))
            continue
        if line.startswith(("考察要点：", "考察要点:", "评价要点：", "评价要点:")) and questions:
            questions[-1]["evaluationPoints"] = _split_points(line.split(":", 1)[-1].split("：", 1)[-1])
            continue
        if (line.endswith(("？", "?")) or len(line) <= 80) and "题单" not in line and "面试" not in line:
            questions.append(_plain_question(line))
    seen: set[str] = set()
    result: list[dict] = []
    for question in questions:
        key = re.sub(r"\s+", "", question["question"])
        if key and key not in seen:
            seen.add(key)
            result.append(question)
    return result


def _plain_question(text: str) -> dict:
    return {
        "question": text.strip(),
        "evaluationPoints": [],
        "required": True,
        "resultType": "capability",
    }
